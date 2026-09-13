"""Referatni Word (.docx) va PDF formatlarida yaratish.

DOCX: python-docx
PDF: reportlab (platypus)
Shriftlar: app/assets/fonts ichidan (Arial) — yo'q bo'lsa Windows fontlari.
"""

import io
import os
import pathlib
import re

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

ASSETS_DIR = pathlib.Path(__file__).resolve().parent / "assets" / "fonts"

_FONT_CACHE: dict[str, str] = {}


def _path(name: str) -> str:
    if name in _FONT_CACHE:
        return _FONT_CACHE[name]
    local = ASSETS_DIR / name
    if local.exists():
        _FONT_CACHE[name] = str(local)
        return str(local)
    win = os.path.join("C:\\Windows\\Fonts", name)
    if os.path.exists(win):
        _FONT_CACHE[name] = win
        return win
    _FONT_CACHE[name] = ""
    return ""


def sanitize_filename(title: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|]+', "_", title).strip()
    return cleaned[:80] or "referat"


def display_name(title: str) -> str:
    return f"{sanitize_filename(title)}.docx"


def _split_blocks(content: str) -> list[str]:
    return [b.strip() for b in content.split("\n") if b.strip()]


def _is_heading(line: str) -> bool:
    return bool(
        re.match(r"^(\d+\.|Kirish|Xulosa|X U L O S A|Foydalanilgan|Adabiyot)", line, re.I)
    )


def build_docx(title: str, content: str) -> bytes:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    h = doc.add_heading(title, level=0)
    for run in h.runs:
        run.font.name = "Arial"

    for block in _split_blocks(content):
        if _is_heading(block):
            heading = doc.add_heading("", level=1)
            run = heading.add_run(block)
            run.font.name = "Arial"
            run.font.size = Pt(13)
        else:
            p = doc.add_paragraph(block)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _register_pdf_fonts() -> bool:
    if "SH" in pdfmetrics.getRegisteredFontNames():
        return True
    regular = _path("arial.ttf")
    bold = _path("arialbd.ttf")
    oblique = _path("ariali.ttf")
    if not regular:
        return False
    fonts = [("SH", regular)]
    if bold:
        fonts.append(("SH-Bold", bold))
    if oblique:
        fonts.append(("SH-Oblique", oblique))
    for name, path in fonts:
        try:
            pdfmetrics.registerFont(TTFont(name, path))
        except Exception:
            pass
    return "SH" in pdfmetrics.getRegisteredFontNames()


def _html_escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_pdf(title: str, content: str) -> bytes:
    _register_pdf_fonts()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=50, rightMargin=50, topMargin=55, bottomMargin=55,
        title=sanitize_filename(title),
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "Stitle", parent=styles["Title"], fontName="SH-Bold",
        fontSize=17, leading=22, alignment=TA_CENTER, textColor=HexColor("#141A29"),
        spaceAfter=18,
    )
    heading_style = ParagraphStyle(
        "Sheading", parent=styles["Heading1"], fontName="SH-Bold",
        fontSize=13, leading=17, spaceBefore=10, spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "Sbody", parent=styles["BodyText"], fontName="SH",
        fontSize=11, leading=16, alignment=TA_JUSTIFY, spaceAfter=6,
    )

    flow = [Paragraph(_html_escape(title), title_style), Spacer(1, 8)]
    for block in _split_blocks(content):
        if _is_heading(block):
            flow.append(Paragraph(_html_escape(block), heading_style))
        else:
            flow.append(Paragraph(_html_escape(block), body_style))

    def on_page(canvas, doc_):
        canvas.saveState()
        canvas.setFont("SH", 8)
        canvas.drawCentredString(A4[0] / 2, 28, str(doc_.page))
        canvas.restoreState()

    doc.build(flow, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()